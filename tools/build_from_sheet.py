"""
Rebuild the project pages and project lists from the content spreadsheet.

Usage (from the Portfolio folder):
    python tools/build_from_sheet.py path/to/Projects.xlsx

Reads the "Main Projects" tab (one column per project, one row per section)
and the "Mini Projects" tab, then writes:
  - projects/<slug>.html          one page per Main Project column with content
  - assets/js/projects.js          card blurbs (Hook row) + More Projects list

Text conventions in the sheet:
  - Blank line                    -> new paragraph
  - A line ending in ":"          -> the lines after it become a bullet list
                                      (in "At a glance", it becomes a small label)
  - "Matlab script:" / "Code:"    -> everything after it is shown as a
                                      collapsible code block
  - [NEEDS INPUT ...] [CONFIRM ...] -> highlighted yellow so drafts are easy to spot
Requires: pip install openpyxl
"""
import html
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent

# Sheet column title -> page slug (must match SELECTED_WORK slugs)
SLUGS = {
    "PearCat": "pearcat",
    "Synthia": "synthia",
    "Cantilevered DeltaXY": "cantilevered-deltaxy",
    "TwoFold": "twofold",
    "Layer by Layer": "layer-by-layer",
    "Small-scale Pool Boiling Facility": "pool-boiling",
    "MochiMash": "mochimash",
    "YubbaYubba": "yubbayubba",
    "en clair": "en-clair",
    "Unit Cell Tessellation": "unit-cell-tessellation",
    "Archetype": "archetype-ai",
    "Discover Mechanical Engineering (DME)": "dme-2026",
}

ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]
TODO_RE = re.compile(r"\[(?:NEEDS INPUT|CONFIRM)[^\]]*\]|^\[[^\]]+\]$", re.M)
CODE_START_RE = re.compile(r"^(matlab script|matlab|code|python script)[^:]*:\s*$", re.I)


def esc(s):
    return html.escape(s, quote=False)


def inline(text):
    """Escape text and highlight [NEEDS INPUT]/[CONFIRM] notes."""
    out, last = [], 0
    for m in TODO_RE.finditer(text):
        out.append(esc(text[last:m.start()]))
        out.append(f'<span class="todo">{esc(m.group(0))}</span>')
        last = m.end()
    out.append(esc(text[last:]))
    return "".join(out)


def is_equation(line):
    return ("=" in line or "³" in line) and len(line.split()) < 14 and not line.rstrip().endswith(".")


def render_body(text, para_class="body-text", split_code=False):
    """Turn a cell into HTML paragraphs / lists / code blocks."""
    if not text:
        return ("", "") if split_code else ""
    lines = text.replace("\r", "").split("\n")
    # split off a trailing code block
    code = None
    for i, ln in enumerate(lines):
        if CODE_START_RE.match(ln.strip()):
            code = "\n".join(lines[i + 1:]).strip("\n")
            label = ln.strip().rstrip(":")
            lines = lines[:i]
            break
    parts, i = [], 0
    while i < len(lines):
        ln = lines[i].strip()
        if not ln:
            i += 1
            continue
        if ln.endswith(":") and i + 1 < len(lines) and lines[i + 1].strip():
            parts.append(f'<p class="{para_class}">{inline(ln)}</p>')
            items = []
            i += 1
            while i < len(lines) and lines[i].strip():
                items.append(f"<li>{inline(lines[i].strip())}</li>")
                i += 1
            parts.append(f'<ul class="body-list">{"".join(items)}</ul>')
            continue
        foot = re.match(r"^(.*=.*?)\s+(\*\S.*)$", ln)
        if foot:  # "formula *footnote"
            parts.append(f'<p class="equation">{inline(foot.group(1))}</p>')
            parts.append(f'<p class="footnote">{inline(foot.group(2))}</p>')
        else:
            cls = "equation" if is_equation(ln) else para_class
            parts.append(f'<p class="{cls}">{inline(ln)}</p>')
        i += 1
    code_html = (f'<details class="code"><summary>View {esc(label)}</summary>'
                 f"<pre><code>{esc(code)}</code></pre></details>") if code else ""
    if split_code:
        return "\n          ".join(parts), code_html
    return "\n          ".join(parts + ([code_html] if code_html else []))


def render_glance(text):
    """'At a glance' panel: blank-line groups; 'Label:' lines become small labels."""
    groups = [g for g in re.split(r"\n\s*\n", text.strip()) if g.strip()]
    out = []
    for g in groups:
        rows = []
        for ln in g.split("\n"):
            ln = ln.strip()
            if not ln:
                continue
            if ln.endswith(":"):
                rows.append(f'<span class="info__label info__label--head">{inline(ln[:-1])}</span>')
            elif ":" in ln and not ln.startswith("[") and ln.index(":") < 24:
                k, v = ln.split(":", 1)
                rows.append(f'<span class="info__label">{inline(k)}</span>{inline(v.strip())}')
            else:
                rows.append(inline(ln))
        out.append("<p>" + "<br>".join(rows) + "</p>")
    return "\n          ".join(out)


def render_outcome(text):
    items = [ln.strip() for ln in text.split("\n") if ln.strip()]
    return '<ul class="outcome-list">' + "".join(f"<li>{inline(x)}</li>" for x in items) + "</ul>"


def page(slug, d):
    title = d["title"]
    sections = []

    if d.get("overview"):
        sections.append(f'''      <section id="overview" data-nav="Project Overview">
        <div class="project__summary">
          {render_body(d["overview"], "summary-text")}
        </div>
      </section>''')

    for key, sid, nav, heading in [
        ("context", "context", "Context/ Motivation", "Context/Motivation"),
        ("framing", "problem", "Problem Framing", "Problem Framing"),
    ]:
        if d.get(key):
            sections.append(f'''      <section id="{sid}" data-nav="{nav}">
        <h2>{heading}</h2>
          {render_body(d[key])}
      </section>''')

    if d["decisions"]:
        blocks = []
        for n, dec in enumerate(d["decisions"]):
            first, _, rest = dec.strip().partition("\n")
            flip = " decision--flip" if n % 2 else ""
            body, code_html = render_body(rest.strip(), split_code=True)
            blocks.append(f'''        <h3>{inline(first.strip())}</h3>
        <div class="decision{flip}">
          <figure class="media"></figure>
          <div class="decision__text">
          {body}
          </div>
        </div>{chr(10) + "        " + code_html if code_html else ""}''')
        sections.append('      <section id="process" data-nav="Process">\n        <h2>Process</h2>\n\n'
                        + "\n\n".join(blocks) + "\n      </section>")

    if d.get("improve"):
        sections.append(f'''      <section id="improvement" data-nav="Room for Improvement">
        <h2>Room for Improvement</h2>
          {render_body(d["improve"])}
      </section>''')

    if d.get("outcome") and d["outcome"].strip().lower() not in ("n/a", "na"):
        sections.append(f'''      <section id="outcome" data-nav="Outcome">
        <h2>Outcome</h2>
        {render_outcome(d["outcome"])}
        <!-- Gallery: put <img> tags inside the figures; rows of 3 / 2 / 4 -->
        <div class="gallery">
          <div class="gallery__row gallery__row--3">
            <figure class="media"></figure><figure class="media"></figure><figure class="media"></figure>
          </div>
          <div class="gallery__row gallery__row--2">
            <figure class="media"></figure><figure class="media"></figure>
          </div>
          <div class="gallery__row gallery__row--4">
            <figure class="media"></figure><figure class="media"></figure><figure class="media"></figure><figure class="media"></figure>
          </div>
        </div>
      </section>''')

    if d.get("reflection"):
        sections.append(f'''      <section id="reflection" data-nav="Reflection">
        <h2>Reflection</h2>
          {render_body(d["reflection"])}
      </section>''')

    desc = re.sub(r"\s+", " ", TODO_RE.sub("", d.get("hook") or title)).strip()
    return f'''<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{esc(title)} — Joshua Rivera Camacho</title>
  <meta name="description" content="{html.escape(desc)}">
  <link rel="stylesheet" href="../assets/css/style.css">
</head>
<!-- Generated from the content spreadsheet by tools/build_from_sheet.py.
     Re-running the script overwrites this file. -->
<body data-page="project" data-root="../" data-slug="{slug}">
  <div class="page page--project">
    <aside class="sidebar" id="sidebar"></aside>

    <main class="project">
      <h1 class="project__title">{esc(title)}</h1>

      <div class="project__hero">
        <figure class="media"><img src="../assets/img/covers/{slug}.jpg" alt="{html.escape(title)}"></figure>
        <div class="project__info">
          {render_glance(d.get("glance") or "")}
        </div>
      </div>

{chr(10).join(sections)}
    </main>
  </div>
  <script src="../assets/js/projects.js"></script>
  <script src="../assets/js/site.js"></script>
</body>
</html>
'''


def cell(ws, r, c):
    v = ws.cell(r, c).value
    return str(v).strip() if v is not None else ""


def read_main(ws):
    # locate rows by their label in column A / B
    rows = {}
    for r in range(1, ws.max_row + 1):
        a, b = cell(ws, r, 1), cell(ws, r, 2)
        if a:
            rows.setdefault(a.strip('"').lower(), r)
        if re.match(r"^Decision \d+$", b):
            rows.setdefault("decisions", []).append(r)
    projects = {}
    for c in range(3, ws.max_column + 1):
        head = cell(ws, 1, c)
        if not head:
            continue
        key = next((k for k in SLUGS if head.lower().startswith(k.lower())), None)
        if not key:
            print(f"  ! no slug for column '{head}' — skipped")
            continue
        get = lambda name: cell(ws, rows[name], c) if name in rows else ""
        d = {
            "title": head,
            "hook": get("hook"),
            "overview": get("overview"),
            "glance": get("at a glance"),
            "context": get("context/motivation"),
            "framing": get("problem framing"),
            "decisions": [cell(ws, r, c) for r in rows.get("decisions", []) if cell(ws, r, c)],
            "improve": get("room for improvement"),
            "outcome": get("outcome"),
            "reflection": get("reflection"),
        }
        projects[SLUGS[key]] = d
    return projects


SEASON_MONTH = {"january": "01", "winter": "01", "spring": "03", "summer": "06", "fall": "09", "autumn": "09"}


def date_key(label):
    m = re.search(r"(january|winter|spring|summer|fall|autumn)\s+(\d{4})", label, re.I)
    return f"{m.group(2)}-{SEASON_MONTH[m.group(1).lower()]}" if m else ""


def read_mini(ws):
    rows = {cell(ws, r, 2).lower(): r for r in range(1, ws.max_row + 1) if cell(ws, r, 2)}
    out = []
    for c in range(3, ws.max_column + 1):
        title = cell(ws, rows["title"], c)
        if not title:
            continue
        label = TODO_RE.sub("", cell(ws, rows["date"], c)).strip()
        out.append({
            "title": TODO_RE.sub("", title).strip(),
            "blurb": "",
            "caption": TODO_RE.sub("", cell(ws, rows["one-liner"], c)).strip(),
            "date": date_key(label),
            "dateLabel": label,
            "effort": 0,
            "image": "",
            "link": "",
        })
    return out


def update_projects_js(main, mini):
    p = ROOT / "assets/js/projects.js"
    js = p.read_text()
    # card blurbs = Hook row
    for slug, d in main.items():
        if not d["hook"]:
            continue
        blurb = json.dumps(re.sub(r"\s+", " ", TODO_RE.sub("", d["hook"])).strip(), ensure_ascii=False)
        pat = re.compile(r'(\{ slug: "%s",.*?blurb: )".*?"(,)' % re.escape(slug), re.S)
        js, n = pat.subn(lambda m: m.group(1) + blurb + m.group(2), js, count=1)
        if not n:
            print(f"  ! couldn't set blurb for {slug}")
    # More Projects list
    items = ",\n".join("  " + json.dumps(x, ensure_ascii=False) for x in mini)
    js, n = re.subn(r"window\.MORE_PROJECTS = \[.*?\n\];",
                    "window.MORE_PROJECTS = [\n" + items + "\n];", js, count=1, flags=re.S)
    if not n:
        print("  ! couldn't find MORE_PROJECTS")
    p.write_text(js)


def main():
    xlsx = sys.argv[1]
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    projects = read_main(wb["Main Projects"])
    for slug, d in projects.items():
        has_content = any(d[k] for k in ("overview", "glance", "context", "framing", "decisions"))
        if not has_content:
            print(f"  - {slug}: no content in sheet, page left as is")
            continue
        (ROOT / "projects" / f"{slug}.html").write_text(page(slug, d))
        print(f"  ✓ projects/{slug}.html")
    mini = read_mini(wb["Mini Projects"]) if "Mini Projects" in wb.sheetnames else []
    update_projects_js(projects, mini)
    print(f"  ✓ assets/js/projects.js ({len(mini)} More Projects)")


if __name__ == "__main__":
    main()
